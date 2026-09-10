"""Extractors module for text parsing and chunking across multi-format uploads.
Supports TXT, MD, PDF, DOCX, PPTX, CSV, XLSX, and PNG/JPG images.
"""

import io
import os
import re
from typing import List, Tuple

import pypdf
import docx
from pptx import Presentation
import openpyxl
import csv
from PIL import Image

try:
    import pytesseract
    HAS_PYTESSERACT = True
except ImportError:
    HAS_PYTESSERACT = False


def extract_text_from_file(file_path: str, filename: str) -> Tuple[str, str]:
    """Extracts raw text content and determines file type."""
    ext = os.path.splitext(filename)[1].lower().strip(".")
    if not ext and file_path:
        ext = os.path.splitext(file_path)[1].lower().strip(".")

    text = ""
    file_type = ext or "txt"

    try:
        if ext in ["txt", "md", "markdown", "log"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

        elif ext == "pdf":
            reader = pypdf.PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted)
            text = "\n\n".join(pages_text)

        elif ext in ["docx", "doc"]:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)

        elif ext in ["pptx", "ppt"]:
            prs = Presentation(file_path)
            slides_text = []
            for slide in prs.slides:
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide_lines:
                    slides_text.append("\n".join(slide_lines))
            text = "\n\n".join(slides_text)

        elif ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheets_text = []
            for sheet in wb.worksheets:
                sheet_lines = [f"--- Sheet: {sheet.title} ---"]
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val) for val in row if val is not None]
                    if row_vals:
                        sheet_lines.append(" | ".join(row_vals))
                sheets_text.append("\n".join(sheet_lines))
            text = "\n\n".join(sheets_text)

        elif ext == "csv":
            lines = []
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    if row:
                        lines.append(" | ".join(row))
            text = "\n".join(lines)

        elif ext in ["png", "jpg", "jpeg", "webp", "bmp"]:
            image = Image.open(file_path)
            if HAS_PYTESSERACT:
                try:
                    text = pytesseract.image_to_string(image)
                except Exception as ocr_err:
                    text = f"[Image File: {filename} - OCR unavailable: {str(ocr_err)}]"
            else:
                text = f"[Image File: {filename} - Image uploaded. Install pytesseract for automatic OCR.]"

        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

    except Exception as e:
        text = f"Error extracting text from file {filename}: {str(e)}"

    return text.strip(), file_type


def chunk_text(text: str, chunk_size: int = 350, overlap: int = 50) -> List[str]:
    """Splits text into paragraph or word-based searchable chunks."""
    if not text:
        return []

    # First split by paragraphs/double newlines
    paragraphs = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]
    chunks = []
    current_chunk = []
    current_words = 0

    for paragraph in paragraphs:
        words = paragraph.split()
        if not words:
            continue

        if current_words + len(words) <= chunk_size:
            current_chunk.append(paragraph)
            current_words += len(words)
        else:
            if current_chunk:
                chunks.append("\n\n".join(current_chunk))
            # If paragraph itself exceeds chunk_size, split by words
            if len(words) > chunk_size:
                for i in range(0, len(words), chunk_size - overlap):
                    chunks.append(" ".join(words[i:i + chunk_size]))
                current_chunk = []
                current_words = 0
            else:
                current_chunk = [paragraph]
                current_words = len(words)

    if current_chunk:
        chunks.append("\n\n".join(current_chunk))

    return chunks
